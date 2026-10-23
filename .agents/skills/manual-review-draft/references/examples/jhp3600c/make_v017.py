"""v0.17: existing Japanese -> black; no JP corpus -> English in green (needs new translation); count volume.

Usage: python3 -I make_v017.py <v016_dir> <out_dir>
"""
import csv
import json
import re
import shutil
import sys
from pathlib import Path

from bs4 import BeautifulSoup

SRC, OUT = Path(sys.argv[1]), Path(sys.argv[2])
if OUT.exists():
    shutil.rmtree(OUT)
shutil.copytree(SRC, OUT)
HTML = OUT / "manual_jhp3600c_jp_en_draft.html"
soup = BeautifulSoup(HTML.read_text(encoding="utf-8"), "html.parser")

US = "JHP-3600C US en（https://ht-doc.readthedocs.io/JHP-3600C/US/en/md/manual_jhp3600c_us_en.html）"
NO_JP = "日文模板、Translation_Memory、线上规格／LCD 表及日规 Web 手册（JA-AD500A-SIL、JBP-1000B-WH、JBP-5000A、JE-1000E-WH、JE-1000F、JE-5000A）均无对应日文"
TM = "Translation_Memory"


def esc(t):
    return t.replace("&", "&amp;").replace('"', "&quot;")


def frag(h):
    return BeautifulSoup(h, "html.parser")


def Z(t):
    return f'<span class="zh-translation" lang="zh-CN">{t}</span>'


def R(t, src=None):
    small = f'<small class="jp-review-source">{src}</small>' if src else ""
    return f'<div class="jp-review" lang="zh-CN"><strong>审阅：</strong>{t}{small}</div>'


def J(t, src):
    return f'<span class="jp-original" data-jp-original-source="{esc(src)}" lang="ja">{t}</span>'


def N(t, kind="term", ref=None):
    r = f' data-reference-jp="{esc(ref)}"' if ref else ""
    return f'<span class="en-new" data-needs-translation="{kind}" lang="en"{r}>{t}</span>'


def P(t):
    return f'<span class="us-difference jp-parameter">{t}</span>'


def NB(tag, t, zh, kind="sentence", ref=None, extra=""):
    r = f' data-reference-jp="{esc(ref)}"' if ref else ""
    return f'<{tag} class="en-new" data-needs-translation="{kind}" lang="en"{r}>{t}{Z(zh)}{extra}</{tag}>'


def JB(tag, t, zh, src, extra=""):
    return f'<{tag} class="jp-original" data-jp-original-source="{esc(src)}" lang="ja">{t}{Z(zh)}{extra}</{tag}>'


def callout(label_jp, zh_label, body):
    return (f'<div class="jp-callout"><strong class="jp-original" data-jp-original-source="{TM}: {label_jp}" lang="ja">'
            f"{label_jp}{Z(zh_label)}</strong>{body}</div>")


# ------------------------------------------------------------ styles + legend
style = soup.new_tag("style")
style.string = (
    "html body #furo-main-content .jp-spec-filled.jp-spec-filled.jp-spec-filled,"
    "html body #furo-main-content :is(.jp-original,.us-difference,.jp-change) span.jp-spec-filled.jp-spec-filled"
    "{color:inherit !important;font-weight:inherit}"
    "html body #furo-main-content .en-new.en-new.en-new,"
    "html body #furo-main-content .en-new.en-new :is(strong,span,a,.jp-spec-filled),"
    "html body #furo-main-content :is(.jp-original,.us-difference,.jp-change,td) .en-new.en-new"
    "{color:#047857 !important}"
    "html body #furo-main-content .en-new .zh-translation.zh-translation{color:#1657bb !important;font-weight:400 !important}"
    "html body #furo-main-content .en-new .jp-parameter.jp-parameter{color:#c2185b !important}"
    "html body #furo-main-content h1.en-new{background:#ecfdf5 !important;border:1px solid #a7f3d0 !important}"
    "html body.only-changes #furo-main-content .en-new{background:#ecfdf5;outline:1px solid #a7f3d0;outline-offset:2px}"
    "html body.only-changes #furo-main-content .jp-spec-filled{background:none;outline:none}"
)
soup.head.append(style)

banner = soup.find(class_="draft-banner")
banner.find("strong").string = "v0.17 · 有日文用日文（黑） · 缺日文语料用英文（绿，需新增翻译） · 中文对照（蓝）"
legend = banner.find(class_="legend")
legend.clear()
legend.append(frag(
    '<span><i class="black"></i>黑色：现成日文（模板／TM／日规 Web 手册）</span>'
    '<span><i class="green"></i>绿色：缺日文语料，暂用英文（需新增翻译）</span>'
    '<span><i class="blue"></i>蓝色小字：中文翻译</span>'
    '<span><i class="rose"></i>玫红：待填参数／待确认</span>'
    '<span><i class="orange"></i>橙色：中文审阅</span>'))
banner.find(class_="legend-detail").string = (
    "日文语料范围：日文模板、Translation_Memory、线上规格参数表与 LCD 图标表、日规 Web 手册 6 本"
    "（JA-AD500A-SIL、JBP-1000B-WH、JBP-5000A、JE-1000E-WH、JE-1000F、JE-5000A）。规格书数值已并入日文（黑色）。"
    "绿色英文的逐条清单与字数见 new-translation-needed.xlsx。")
soup.find("header").find("small").string = "本地审阅 · v0.17 · 有日文用日文＋缺语料英文（绿）＋新增翻译量统计"
soup.title.string = "JHP-3600C · 日规说明书审核稿 v0.17"

# ------------------------------------------------------------ green-reference texts that meant "spec value"
for old, new in [
    ("v0.15 已按日规规格书回填（绿色）", "v0.15 已按日规规格书回填（数值并入日文，黑色）"),
    ("v0.15 按日规规格书回填（绿色）。", "v0.15 按日规规格书回填（数值并入日文，黑色）。"),
    ("绿色为 2026-10-12 规格书数值；玫红为规格书未列、待确认项。", "数值来自 2026-10-12 规格书；玫红为规格书未列、待确认项。"),
]:
    hit = soup.find(string=lambda s, old=old: s and old in s)
    assert hit is not None, old
    hit.replace_with(hit.replace(old, new))

# ------------------------------------------------------------ existing English units -> green
for el in soup.select("[data-unmatched-reason]"):
    if el.find_parent(id="connections"):
        continue
    cls = [c for c in (el.get("class") or []) if c != "us-difference"]
    el["class"] = cls + ["en-new"]
    words = len(re.findall(r"[A-Za-z]+", el.find(string=True, recursive=False) or ""))
    el["data-needs-translation"] = "term" if words <= 5 else "sentence"
    reason = el["data-unmatched-reason"]
    el["data-unmatched-reason"] = reason if "RTD" in reason or "日规 Web" in reason else reason + "（另查日规 Web 手册 6 本亦无）"
    td = el.find_parent("td")
    if td is not None and "us-difference" in (td.get("class") or []):
        td["class"] = [c for c in td["class"] if c != "us-difference"]

# spec value: English port term inside Japanese value
hit = soup.find(string=lambda s: s and "（Parallel Communication ×1）" in s)
before, after = hit.split("（Parallel Communication ×1）")
hit.replace_with(frag(before + "（" + N("Parallel Communication") + " ×1）" + after))

# ------------------------------------------------------------ LCD: new icon descriptions (US English, green)
LCD_DESC = {
    "Parallel": ("On: The parallel connection is set up successfully.<br/>Blink: The parallel connection is setting up.<br/>Off: The parallel connection is not set up.",
                 "常亮：并机连接已成功建立。闪烁：正在建立并机连接。熄灭：未建立并机连接。"),
    "AC Expansion Indicator (Charging)": ("The product can be charged from the grid through a Jackery Automatic Transfer Switch (ATS).",
                                          "可通过 Jackery 自动切换开关（ATS）从电网为产品充电。"),
    "AC Expansion Indicator (Discharging)": ("The product supplies power to your home loads through the connected ATS.",
                                             "产品通过已连接的 ATS 为家庭负载供电。"),
    "Transfer Switch Indicator": ("On: The product is successfully connected to a transfer switch (ATS).<br/>Off: The product is not connected to a transfer switch (ATS).",
                                  "常亮：产品已成功连接切换开关（ATS）。熄灭：产品未连接切换开关（ATS）。"),
    "EV Charging": ("The Ground Adapter (sold separately) is connected to the product successfully.",
                    "接地适配器（需另购）已成功连接到产品。"),
    "Smart Meter indicator": ("On: A smart meter is online.<br/>Off: No smart meter is added to the system, or the smart meter is offline.",
                              "常亮：智能电表在线。熄灭：系统未添加智能电表，或智能电表离线。"),
}
lcd = soup.find(id="lcd-display")
filled_lcd = []
for ph in lcd.select('p[data-editorial-placeholder="unfilled"]'):
    tr = ph.find_parent("tr")
    names = [e.find(string=True, recursive=False).strip() for e in tr.select(".en-new") if e.find(string=True, recursive=False)]
    key = next((n for n in names if n in LCD_DESC), None)
    if key is None:  # merged description row (Charging/Discharging share one row group)
        prev = tr.find_previous_sibling("tr")
        names = [e.find(string=True, recursive=False).strip() for e in prev.select(".en-new")] if prev else []
        key = next((n for n in names if n in LCD_DESC and n not in filled_lcd), None)
    assert key, names
    en, zh = LCD_DESC[key]
    ph.replace_with(frag(NB("p", en, zh, kind="sentence")))
    filled_lcd.append(key)
    note = tr.find(string=lambda s: s and "现有日文 LCD 句库未找到此项对应原句" in s)
    if note is not None:
        note.replace_with("日文 LCD 语料及日规 Web 手册均无此图标说明，暂用美规英文（绿色），需新增翻译。")
missing = [k for k in LCD_DESC if k not in filled_lcd]
for key in missing:  # icon rows whose description cell is not a placeholder: append under the name
    name_el = next(e for e in lcd.select(".en-new") if (e.find(string=True, recursive=False) or "").strip() == key)
    desc_td = name_el.find_parent("tr").find_all("td")[-1]
    en, zh = LCD_DESC[key]
    desc_td.append(frag(NB("p", en, zh, kind="sentence")))
    filled_lcd.append(key)

# ------------------------------------------------------------ CONNECTIONS chapter rebuilt
old = soup.find(id="connections")
figs = {f["data-us-image-source"].rsplit("#", 1)[1]: str(f) for f in old.find_all("figure")}
PN = P("[PRODUCT NAME]")
SRC_CONNECT = "日规 Web 手册见出し「接続」（JA-AD500A-SIL、JE-1000E-WH）"
SRC_ACC = "docs/templates/page_jp/08_charging_methods.rst「必要なアクセサリー」"
SRC_SOLD = TM + ": CONNECT TO BATTERY PACK(S) (SOLD SEPARATELY) → バッテリーパックに接続する（別売り）"
SRC_SOLD2 = TM + ": The car charging cable and Jackery DC Input Module are sold separately. → …は別売りです。"
SRC_REF = "日规 Web 手册 JE-5000A／JE-1000E-WH：使用方法の詳細については、…の取扱説明書を参照してください。"
SRC_OUTAGE = "日规 Web 手册 JE-1000F／JE-1000E-WH：停電時には家庭用電源からポータブル電源の切り替え時間が10ミリ秒で行われます。"
REF_OFF = "…に接続する前に、両方の電源がオフになっていることを確認してください。（JE-5000A／JBP-5000A，仅覆盖“两台均关机”）"
REF_LOAD = "複数の機器をUPSモードで同時に接続する場合は、合計負荷電力が…を超えないようにしてください。（06_ups_mode.rst／JE-1000F）"
REF_OVER = "超えると、機器が過負荷保護により停止するおそれがあります。（06_ups_mode.rst／JE-1000F）"
REF_HOME = "高出力機器に電力を供給し、家庭用電源のインレットボックスまたは切替スイッチに接続します。（JE-5000A）"

chapter = (
    '<section data-english-by-request="2026-10-12" data-jp-layout-source="JHP-3600C US en · CONNECTIONS" id="connections">'
    + JB("h1", "接続", "连接", SRC_CONNECT)
    + '<div class="jp-change">'
    + R("本章按用户指示新增（美规 CONNECTIONS 章按日规规格书改写，EPO 不纳入；电池包连接仍在「充電方法」章）。"
        "已有日文的（見出し「接続」「必要なアクセサリー」「ATS／MTSに接続する（別売り）」、「…は別売りです。」、注意／説明、"
        "参照取扱説明書句、停電切替時間句）已换成现成日文（黑色）；其余在日文语料中找不到，暂用英文（绿色），计入新增翻译量。"
        "[PRODUCT NAME] 为日文商品名占位。", "版面依据：JHP-3600C US en · CONNECTIONS")
    # cascade
    + NB("h2", "CASCADE PARALLEL CONNECTION", "级联并机", kind="heading")
    + NB("p", f"Cascade connection allows 2 {PN} units to operate as a combined system, increasing total output power. "
              "Total output: up to 6000W; total capacity with battery packs: up to 43kWh.",
         "级联连接可使 2 台本产品作为一个系统运行，提高总输出功率。总输出最大 6000W；连接电池包后总容量最大 43kWh。")
    + JB("h3", "必要なアクセサリー", "所需配件", SRC_ACC)
    + JB("p", N("Jackery 30A Parallel/Fast Charging Cable", kind="name") + "および"
              + N("Jackery 30A Parallel Communication Cable", kind="name") + "は別売りです。",
         "Jackery 30A 并机／快充线和 Jackery 30A 并机通讯线需另购。", SRC_SOLD2,
         extra=R("线缆名称按规格书选配件“30A 并机/快充线”“30A 并机通讯线”暂拟英文，正式名称待定；句式为现成日文。"))
    + NB("h3", "Connection Steps", "连接步骤", kind="heading")
    + "<ol>"
    + NB("li", "Ensure both units are powered OFF and completely disconnected from power sources.",
         "确认两台设备均已关机，并与所有电源完全断开。", ref=REF_OFF)
    + NB("li", "Connect Parallel Communication ports on both units.", "连接两台设备的并机通讯口。")
    + "</ol>"
    + callout("注意", "注意", NB("p", "Do not skip this step. Without communication, the units cannot synchronize data and may be damaged.",
                                 "请勿跳过此步骤。未建立通讯时，两台设备无法同步数据，可能导致损坏。"))
    + '<ol start="3">'
    + NB("li", "Connect the first unit's 100V/200V Output Port (NEMA L14-30R) to the second unit's AC Expansion Port.",
         "将第一台的 100V/200V 输出口（NEMA L14-30R）连接到第二台的 AC Expansion Port（并机口）。",
         extra=R("美规为第一台 240V 输出（NEMA 14-50R）接第二台 AC Expansion Port。日规端口按规格书推定为 L14-30R → 并机口，使用 30A 并机/快充线，需确认。"))
    + NB("li", "System information will synchronize across both displays.", "两台显示屏上的系统信息将同步。")
    + "</ol>"
    + figs["cascade"]
    + callout("注意", "注意", NB("p", f"After cascading, the total system output power via the {P('[OUTPUT PORT TBC]')} of the second unit is 6000W. "
                                      "The total load MUST NOT exceed total power.",
                                 "级联后，经第二台【输出口待确认】输出的系统总功率为 6000W。总负载不得超过总功率。", ref=REF_LOAD))
    + R("规格书：级联并机 6kW。L14-30R 规格书为 20A／4000W，6000W 经哪个端口输出待确认。")
    # backup power
    + NB("h2", "BACKUP POWER CONNECTION", "备用电源连接", kind="heading")
    + JB("h3", "ATSに接続する（別売り）", "连接 ATS（需另购）", SRC_SOLD)
    + '<p lang="ja">'
    + N(f"{PN} can supply backup power to home circuits through a Jackery Automatic Transfer Switch (ATS).", kind="sentence", ref=REF_HOME)
    + J("使用方法の詳細については、Jackery ATSの取扱説明書を参照してください。", SRC_REF)
    + Z("本产品可通过 Jackery 自动切换开关（ATS）为家庭回路提供备用电源。使用方法详情请参阅 Jackery ATS 使用说明书。")
    + "</p>"
    + figs["ats-lock"]
    + callout("説明", "说明",
              '<p lang="ja">'
              + N("When UPS mode is enabled on the ATS, the power station remains active and continuously consumes power.", kind="sentence")
              + J("停電時には家庭用電源からポータブル電源の切り替え時間が", SRC_OUTAGE) + P("20ミリ秒【要確認】") + J("で行われます。", SRC_OUTAGE)
              + Z("在 ATS 上启用 UPS 模式时，储能电源保持运行并持续耗电。停电时，家用电源切换到便携式电源的时间为【20 毫秒，待确认】。")
              + "</p>")
    + R("美规 ATS 模式切换为 20 ms；规格书只写本机 UPS 10ms，ATS 场景切换时间待确认。日规 ATS 型号待确认。")
    + JB("h3", "MTSに接続する（別売り）", "连接 MTS（需另购）", SRC_SOLD)
    + NB("p", f"{PN} can be connected to a Manual Transfer Switch (MTS) to power selected home circuits. "
              "Choose the appropriate method based on your installation mode.",
         "本产品可连接手动切换开关（MTS），为选定的家庭回路供电。请根据安装方式选择合适的连接方法。")
    + callout("説明", "说明", NB("p", "All cables used in this section are sold separately or included with the separately sold MTS.",
                                 "本节使用的所有线缆均需另购，或随另购的 MTS 附带。"))
    + callout("注意", "注意", NB("p", "The built-in inverter of this product is rated at 3.6kW. When connecting to an MTS, pay attention to the total load power. "
                                      "A load exceeding the inverter power may cause an overload shutdown.",
                                 "由于主机逆变器只有 3.6kW，接入 MTS 时需要注意负载功率，大于逆变器功率可能会过载断电。", ref=REF_OVER))
    + R("上述注意事项按日规规格书备注新增（英文为译写）。美规 GFCI 安装提示属美规电气体系，未写入正文，适用性待电气确认。")
    + NB("h3", "SINGLE UNIT CONNECTION WITH MTS", "单机连接 MTS", kind="heading")
    + NB("p", f"In automatic mode, {PN} receives AC input and provides UPS-like backup power through the MTS.",
         "自动模式下，本产品接收交流输入，并通过 MTS 提供类似 UPS 的备用电源。")
    + "<ol>"
    + NB("li", f"Power OFF {PN}.", "关闭本产品电源。")
    + NB("li", f"Connect {PN} to a {P('[100V / 200V TBC]')} AC power supply.", "将本产品连接到【100V／200V 待确认】交流电源。",
         extra=R("日规可经 AC 充电口（100V，1500W）或经并机口用快充线接 200V，规格书未说明，需确认输入方式及旁路功率。"))
    + NB("li", "Connect the 100V/200V Output Port (NEMA L14-30R) to the MTS inlet.", "将 100V/200V 输出口（NEMA L14-30R）连接到 MTS 入口。",
         ref=REF_HOME)
    + NB("li", "Turn the load switch of the MTS to backup.", "将 MTS 的负载开关拨到备用（backup）。")
    + NB("li", "Turn ON the unit and enable AC Output.", "开启本产品并打开交流输出。")
    + "</ol>"
    + figs["mts-single"]
    + '<p lang="ja">'
    + N("When grid power is present, AC power passes through to the MTS.", kind="sentence")
    + J("停電時には家庭用電源からポータブル電源の切り替え時間が10ミリ秒で行われます。", SRC_OUTAGE)
    + Z("有市电时，交流电直通至 MTS。停电时，家用电源切换到便携式电源的时间为 10 毫秒。") + "</p>"
    + callout("説明", "说明", NB("p", "Automatic switching requires the AC input to remain connected at all times.", "自动切换要求交流输入始终保持连接。"))
    + NB("h3", "CASCADE PARALLEL CONNECTION WITH MTS", "级联并机连接 MTS", kind="heading")
    + NB("p", "When two units are connected in cascade parallel mode, the system can deliver higher output power to the MTS.",
         "两台设备以级联并机方式连接时，系统可向 MTS 提供更高的输出功率。")
    + "<ol>"
    + NB("li", "Complete the Cascade Parallel Connection steps described earlier.", "先完成前述级联并机连接步骤。")
    + NB("li", "Connect the second unit's 100V/200V Output Port (NEMA L14-30R) to the MTS inlet.", "将第二台的 100V/200V 输出口（NEMA L14-30R）连接到 MTS 入口。")
    + NB("li", "Turn ON both units and enable AC Output.", "开启两台设备并打开交流输出。")
    + "</ol>"
    + figs["mts-cascade"]
    + callout("注意", "注意", NB("p", "After cascading, the maximum output to the MTS is 6000W. The total load MUST NOT exceed total power.",
                                 "级联后，输出至 MTS 的最大功率为 6000W。总负载不得超过总功率。", ref=REF_LOAD))
    + R("规格书级联 6kW；L14-30R 规格书为 20A／4000W，经 L14-30R 输出 6000W 是否成立需确认。")
    + "</div></section>")
old.replace_with(frag(chapter))
soup.select_one('aside.toc a[href="#connections"]').string = "连接"

# ------------------------------------------------------------ count
SECTION_ZH = {}
for a in soup.select("aside.toc a"):
    SECTION_ZH[a["href"][1:]] = a.get_text(strip=True)


def clean_text(el):
    e = BeautifulSoup(str(el), "html.parser")
    for z in e.select(".zh-translation, .jp-review, ol, figure"):
        z.decompose()
    t = e.get_text(" ", strip=True)
    t = re.sub(r"\s+([.,;:)])", r"\1", re.sub(r"\s+", " ", t)).strip()
    return t


def zh_text(z):
    return re.sub(r"\s+", " ", z.get_text(" ", strip=True)).strip() if z is not None else ""


rows = []
for el in soup.select("[data-needs-translation]"):
    sec = el.find_parent("section")
    while sec.find_parent("section"):
        sec = sec.find_parent("section")
    text = clean_text(el)
    zh_el = el.find(class_="zh-translation", recursive=False) or (el.parent.find(class_="zh-translation", recursive=False) if el.name == "span" else None)
    words = re.findall(r"[A-Za-z0-9][A-Za-z0-9®'’/+\-.]*", re.sub(r"\[[^\]]*\]", " ", text))
    sents = [s for s in re.split(r"(?<=[.!?])\s+", text) if re.search(r"[A-Za-z]", s)]
    rows.append({
        "章节": SECTION_ZH.get(sec.get("id"), sec.get("id")),
        "类型": {"heading": "标题", "term": "术语／名称", "name": "配件名称", "sentence": "句子"}[el["data-needs-translation"]],
        "英文原文": text,
        "英文词数": len(words),
        "句数": len(sents) if el["data-needs-translation"] == "sentence" else 0,
        "英文字符数": len(text),
        "中文参考": zh_text(zh_el),
        "可参考的现成日文（近似，非直接复用）": el.get("data-reference-jp", ""),
        "备注": el.get("data-unmatched-reason", "") or ("按用户要求新增的连接章节" if el.find_parent(id="connections") else ""),
    })
uniq = {}
for r in rows:
    uniq.setdefault(r["英文原文"], r)
summary = {
    "units": len(rows), "unique_units": len(uniq),
    "words": sum(r["英文词数"] for r in rows), "unique_words": sum(r["英文词数"] for r in uniq.values()),
    "chars": sum(r["英文字符数"] for r in rows),
    "by_type": {}, "by_chapter": {},
}
for r in rows:
    for key, k in (("by_type", r["类型"]), ("by_chapter", r["章节"])):
        d = summary[key].setdefault(k, {"units": 0, "words": 0})
        d["units"] += 1
        d["words"] += r["英文词数"]
pending = len(soup.select(".jp-parameter, [data-editorial-placeholder]"))
summary["pending_parameters"] = pending
jp_units = len(soup.select("#furo-main-content .jp-original"))
summary["jp_original_elements"] = jp_units

# CSV (Excel-friendly) + XLSX
fields = list(rows[0].keys())
with open(OUT / "new-translation-needed.csv", "w", encoding="utf-8-sig", newline="") as f:
    w = csv.DictWriter(f, fieldnames=["序号"] + fields)
    w.writeheader()
    for i, r in enumerate(rows, 1):
        w.writerow({"序号": i, **r})
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
wb = Workbook()
ws = wb.active
ws.title = "新增翻译清单"
ws.append(["序号"] + fields)
for i, r in enumerate(rows, 1):
    ws.append([i] + [r[k] for k in fields])
for c in ws[1]:
    c.font = Font(bold=True)
    c.fill = PatternFill("solid", fgColor="E8F5E9")
for col, width in zip("ABCDEFGHIJ", (6, 14, 10, 70, 10, 8, 10, 50, 60, 50), strict=False):
    ws.column_dimensions[col].width = width
for row in ws.iter_rows(min_row=2):
    for c in row:
        c.alignment = Alignment(wrap_text=True, vertical="top")
ws.freeze_panes = "A2"
s2 = wb.create_sheet("汇总")
s2.append(["指标", "数值"])
for k, v in (("需新增翻译单元（总）", summary["units"]), ("去重后单元", summary["unique_units"]),
             ("英文词数（总）", summary["words"]), ("英文词数（去重）", summary["unique_words"]),
             ("英文字符数（总）", summary["chars"]), ("待填参数（非翻译）", pending)):
    s2.append([k, v])
s2.append([])
s2.append(["按类型", "单元", "词数"])
for k, d in summary["by_type"].items():
    s2.append([k, d["units"], d["words"]])
s2.append([])
s2.append(["按章节", "单元", "词数"])
for k, d in summary["by_chapter"].items():
    s2.append([k, d["units"], d["words"]])
s2.column_dimensions["A"].width = 26
for c in s2[1]:
    c.font = Font(bold=True)
wb.save(OUT / "new-translation-needed.xlsx")

# appendix summary block
app = soup.find(id="jp-review-appendix")
chap_rows = "".join(f"<tr><td>{k}</td><td>{d['units']}</td><td>{d['words']}</td></tr>" for k, d in summary["by_chapter"].items())
type_rows = "".join(f"<tr><td>{k}</td><td>{d['units']}</td><td>{d['words']}</td></tr>" for k, d in summary["by_type"].items())
app.find("h1").insert_after(frag(
    "<h2>需新增翻译量（v0.17，绿色英文）</h2>"
    f"<p>共 <strong>{summary['units']}</strong> 个单元（去重 {summary['unique_units']}），英文 <strong>{summary['words']}</strong> 词"
    f"（去重 {summary['unique_words']} 词），{summary['chars']} 字符。待填参数 {pending} 处属于数值确认，不计入翻译。"
    "图中文字（美规原图标注）不计入，重绘时另行处理。逐条清单见 new-translation-needed.xlsx。</p>"
    '<div class="jp-table-scroll"><table class="jp-table"><thead><tr><th>章节</th><th>单元</th><th>词数</th></tr></thead>'
    f"<tbody>{chap_rows}</tbody></table></div>"
    '<div class="jp-table-scroll"><table class="jp-table"><thead><tr><th>类型</th><th>单元</th><th>词数</th></tr></thead>'
    f"<tbody>{type_rows}</tbody></table></div>"))
cl = soup.find(class_="change-log")
cl.find("summary").insert_after(frag(
    "<p><strong>v0.17（2026-10-12）</strong> — 配色改为：现成日文黑色、缺日文语料的英文绿色（需新增翻译）、中文蓝色小字；"
    "连接章节中已有日文的标题、句式与句子换成现成日文；LCD 新增图标补美规英文说明；输出新增翻译清单与字数统计。</p>"))

HTML.write_text(str(soup), encoding="utf-8")
(OUT / "new-translation-summary.json").write_text(json.dumps(summary, ensure_ascii=False, indent=1), encoding="utf-8")
print(json.dumps(summary, ensure_ascii=False, indent=1))
