"""Merge per-group US↔JP alignment results into a coverage summary, an xlsx and a json.

  python3 -I aggregate_align.py <align_dir> <out.xlsx> <out.json>

<align_dir> holds:
  us_units.json / jp_units.json      extract_units.py output (US may include extra image-label units)
  <group>.json + <group>_result.json one pair per alignment agent
  duplicates.json   (optional)       US unit ids to ignore, e.g. a spec table repeated in an annex
  overrides.json    (optional)       {"US123": ["covered", ["JP045"], "reason"]} for cross-group links
                                     and units later fixed in the draft
Result schema per group: {"us": [{"id", "status": covered|partial|missing|us_only, "jp_ids", "note"}],
                          "jp_only": [{"id", "note"}]}
"""
from __future__ import annotations

import json
import sys
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

ZH = {"covered": "已对应", "partial": "部分对应", "missing": "缺失", "us_only": "美规独有（日规不适用）"}
LABELS = {"注意", "説明", "警告", "危険", "ご注意", "重要", "ヒント"}
FILL = {"partial": "FFF7ED", "missing": "FDE2E2", "us_only": "EEEEEE"}


def load(path: Path, default=None):
    return json.loads(path.read_text(encoding="utf-8")) if path.exists() else default


def main() -> int:
    d, out_xlsx, out_json = Path(sys.argv[1]), sys.argv[2], sys.argv[3]
    us = {u["id"]: u for u in load(d / "us_units.json")}
    jp = {u["id"]: u for u in load(d / "jp_units.json")}
    dup = set(load(d / "duplicates.json", []))
    res: dict[str, dict] = {}
    jp_only: dict[str, dict] = {}
    problems = []
    for rf in sorted(d.glob("*_result.json")):
        g = rf.name[: -len("_result.json")]
        r = load(rf)
        expected = {u["id"] for u in load(d / f"{g}.json")["us"]}
        got = [x["id"] for x in r["us"]]
        if set(got) != expected or len(got) != len(set(got)):
            problems.append([g, sorted(expected - set(got))[:5], sorted(set(got) - expected)[:5]])
        for x in r["us"]:
            if x["status"] not in ZH:
                problems.append([g, "bad status", x["id"]])
            res[x["id"]] = dict(x, group=g)
        for x in r.get("jp_only", []):
            jp_only.setdefault(x["id"], x)
    for uid, (status, ids, note) in load(d / "overrides.json", {}).items():
        res[uid] = dict(res.get(uid, {"id": uid}), status=status, jp_ids=ids, note=note + "（汇总校正）")
    unaligned = [uid for uid in us if uid not in res and uid not in dup]
    if unaligned:
        problems.append(["no group", unaligned[:8], len(unaligned)])
    mapped = {j for x in res.values() for j in x.get("jp_ids", [])}
    jp_only = {k: v for k, v in jp_only.items() if k not in mapped and jp[k]["text"] not in LABELS}

    chapters: dict[str, dict] = {}
    for uid, u in us.items():
        if uid in dup or uid not in res:
            continue
        c = chapters.setdefault(u["chapter"], dict.fromkeys(("total", "covered", "partial", "missing", "us_only"), 0))
        c["total"] += 1
        c[res[uid]["status"]] += 1
    tot = {k: sum(c[k] for c in chapters.values()) for k in ("total", "covered", "partial", "missing", "us_only")}
    applicable = tot["total"] - tot["us_only"] or 1
    summary = {"chapters": chapters, "total": tot, "applicable": applicable,
               "coverage_full_pct": round(100 * tot["covered"] / applicable, 1),
               "coverage_any_pct": round(100 * (tot["covered"] + tot["partial"]) / applicable, 1),
               "jp_only": len(jp_only), "duplicates_excluded": len(dup), "problems": problems}

    wb = Workbook()
    ws = wb.active
    ws.title = "汇总"
    ws.append(["美规章节", "单元数", "已对应", "部分对应", "缺失", "美规独有", "对应率（含部分）"])
    for name, c in chapters.items():
        app = c["total"] - c["us_only"]
        ws.append([name, c["total"], c["covered"], c["partial"], c["missing"], c["us_only"],
                   f"{round(100 * (c['covered'] + c['partial']) / app)}%" if app else "—"])
    ws.append(["合计", tot["total"], tot["covered"], tot["partial"], tot["missing"], tot["us_only"],
               f"{summary['coverage_any_pct']}%"])
    ws.column_dimensions["A"].width = 52
    w2 = wb.create_sheet("美规逐条对照")
    w2.append(["美规ID", "美规章节", "小节", "美规原文", "状态", "对应日规单元", "日规原文", "日规中文", "说明"])
    for uid, u in us.items():
        if uid in dup or uid not in res:
            continue
        r = res[uid]
        ids = [i for i in r.get("jp_ids", []) if i in jp]
        w2.append([uid, u["chapter"], u.get("section") or "", u["text"], ZH[r["status"]], ", ".join(ids),
                   "\n".join(jp[i]["text"] for i in ids)[:1500], "\n".join(jp[i]["zh"] for i in ids)[:1500],
                   r.get("note", "")])
        if r["status"] in FILL:
            for cell in w2[w2.max_row]:
                cell.fill = PatternFill("solid", fgColor=FILL[r["status"]])
    w3 = wb.create_sheet("日规独有")
    w3.append(["日规ID", "日规章节", "小节", "日规原文", "中文", "说明"])
    for jid, x in jp_only.items():
        j = jp[jid]
        w3.append([jid, j["chapter"], j.get("section") or "", j["text"], j["zh"], x.get("note", "")])
    for sheet, widths in ((w2, (8, 22, 22, 60, 12, 14, 50, 40, 40)), (w3, (8, 18, 22, 60, 50, 40))):
        for col, width in zip("ABCDEFGHI", widths, strict=False):
            sheet.column_dimensions[col].width = width
        for row in sheet.iter_rows(min_row=2):
            for cell in row:
                cell.alignment = Alignment(wrap_text=True, vertical="top")
        sheet.freeze_panes = "A2"
    for sheet in (ws, w2, w3):
        for cell in sheet[1]:
            cell.font = Font(bold=True)
    wb.save(out_xlsx)

    def rows(status):
        return [dict(id=k, chapter=us[k]["chapter"], text=us[k]["text"], note=v.get("note", ""))
                for k, v in res.items() if v["status"] == status and k not in dup]
    Path(out_json).write_text(json.dumps({
        "summary": summary, "missing": rows("missing"), "partial": rows("partial"), "us_only": rows("us_only"),
        "jp_only": [dict(id=k, chapter=jp[k]["chapter"], text=jp[k]["text"], zh=jp[k]["zh"], note=v.get("note", ""))
                    for k, v in jp_only.items()]}, ensure_ascii=False, indent=1), encoding="utf-8")
    print(json.dumps({k: summary[k] for k in ("total", "applicable", "coverage_full_pct", "coverage_any_pct",
                                               "jp_only", "problems")}, ensure_ascii=False))
    return 1 if problems else 0


if __name__ == "__main__":
    sys.exit(main())
